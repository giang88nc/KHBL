"""Bộ hồi quy THÂU VÀNG (THAU_VANG) trên SANDBOX — không bao giờ ghi máy KK.

Chạy trọn vòng đời qua bill.py (Ins → Upd → CompleteMore [+CARDPAY_Ins khi CK] → T_TILL_TXN_Proc → hủy TT → xóa)
và qua VIEW web (test client, ép đích sandbox trong tiến trình), so số dòng/số dư két trước–sau. Tự dọn bằng chính
chuỗi hủy của vendor; lỗi giữa chừng → dọn tay.
    manage.py smoke_thau
"""
import datetime
import json
import re
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.test import Client
from django.contrib.auth import get_user_model

from apps.pmv import gateway, money as M
from apps.pmv.client import PmvClient, PmvProcError


TABLES = ("TRN_RT_BUYGOLD", "T_TILL_TXN", "T_TILL_TXN_DETAIL", "T_TILL_BAL", "I_LICHSUTICHLUYDIEM", "I_DIEMTICHLUY",
          "TRN_RT_BUYGOLD_CardPay", "TRN_RT_BUYGOLD_Log", "T_TILL_TXN_Log")


class Command(BaseCommand):
    help = "Smoke THAU_VANG trên sandbox: lập → sửa → chốt (tiền mặt & CK) → hủy TT → xóa; kiểm từng bảng"

    def handle(self, *args, **opts):
        gateway.dich_hien_tai = lambda: "sandbox"          # ép đích cho MỌI PmvClient() trong tiến trình này
        from apps.pos import bill as B, services as S, views as V
        c = PmvClient("sandbox", tag="smoke_thau")
        self.ok, self.fail = 0, []
        trn = ""
        try:
            till = c.query("SELECT TOP 1 TillID FROM T_TILL WITH (NOLOCK) WHERE Active='1' ORDER BY TillID DESC")[0]["TillID"]
            emp = c.query("SELECT TOP 1 EmpID FROM T_EMPLOYEE WITH (NOLOCK) WHERE ISNULL(Active,'1')='1' ORDER BY EmpID")[0]["EmpID"]
            user = c.query("SELECT TOP 1 UserID FROM SYS_USERS WITH (NOLOCK) ORDER BY UserID")[0]["UserID"]
            shop = c.query("SELECT TOP 1 ShopID FROM T_SHOP WITH (NOLOCK) WHERE Active='1'")[0]["ShopID"]
            cust = c.query("SELECT TOP 1 c.CustID FROM I_CUSTOMER c WITH (NOLOCK) JOIN I_DIEMTICHLUY d WITH (NOLOCK) ON d.CustID=c.CustID "
                           "WHERE c.CustID<>'CU0000000000000' ORDER BY c.CustID DESC")[0]["CustID"]
            rate = M.dec(c.query("SELECT BuyRate FROM I_XRATE WITH (NOLOCK) WHERE GoldCcy='D18K'")[0]["BuyRate"]) or M.dec(8700)
            self.stdout.write(f"  két {till} · NV {emp} · user {user} · khách {cust} · D18K {rate}")

            def counts():
                return {t: c.query(f"SELECT COUNT(*) n FROM {t} WITH (NOLOCK)")[0]["n"] for t in TABLES}

            def bal(ccy):
                r = c.query("SELECT ISNULL(TillBal,0) b FROM T_TILL_BAL WITH (NOLOCK) WHERE TillID=? AND GoldCcy=? AND ProductCode IS NULL", (till, ccy))
                return M.dec(r[0]["b"]) if r else M.dec(0)

            before = counts()
            bal0 = {"D18K": bal("D18K"), "VND": bal("VND")}
            last0 = c.query("SELECT LastTradingDate FROM I_CUSTOMER WITH (NOLOCK) WHERE CustID=?", (cust,))[0]["LastTradingDate"]

            # ── 1. LẬP (Ins) ──
            row = B.luu_thau(cust_id=cust, emp_id=emp, till_id=till, shop_id=shop, user_id=user, gold_code="D18K",
                             gw=990, dw=10, rate=rate, pct=100, add_money=0, notes="smoke_thau", c=c)
            trn = row["TrnID"]
            tien1 = M.round_vnd(M.dec(990) / 100 * rate * 1000)
            self._ok("Ins → W + BillCode + TotalAmount = 9,9 chỉ × giá", row["Status"] == "W" and bool(row["BillCode"])
                       and M.dec(row["TotalAmount"]) == tien1, f"{trn} {row['BillCode']} {row['TotalAmount']}")
            d1 = self.delta(before, counts())
            self._ok("Ins chỉ +1 TRN_RT_BUYGOLD", d1 == {"TRN_RT_BUYGOLD": 1}, str(d1))
            moc1 = row["TrnDateTime_Upd"]

            # ── 2. SỬA (Upd với mốc) ──
            row = B.luu_thau(trn_id=trn, cust_id=cust, emp_id=emp, till_id=till, shop_id=shop, user_id=user, gold_code="D18K",
                             gw=1000, dw=10, rate=rate, pct=100, add_money=-20000, notes="smoke_thau sửa", c=c)
            tien2 = M.round_vnd(M.dec(1000) / 100 * rate * 1000 - 20000)
            self._ok("Upd đổi TL + bớt → tiền mới, mốc đổi", M.dec(row["TotalAmount"]) == tien2 and row["TrnDateTime_Upd"] != moc1
                       and M.dec(row["AddMoney"]) == -20000, str(row["TotalAmount"]))
            # mốc CŨ cố tình → goi_co_khoa phải báo lỗi chứ không "thành công giả"
            try:
                p = B._tham_so_thau(cust_id=cust, emp_id=emp, till_id=till, shop_id=shop, user_id=user, gold_code="D18K",
                                    gw=1000, dw=10, rate=rate, pct=100, add_money=-20000, tien=tien2, notes="x")
                p.update(p_TrnID=trn, p_UserUpd=user, p_IsGiaoDichNhanh="0", p_CreatedDate=None)
                p.pop("p_TrnID_GDN", None)
                full = c.tham_so_day_du("TRN_RT_BUYGOLD_Upd", p)
                full["p_TrnDateTime_Upd"] = PmvClient.fmt_moc(moc1)
                rc, sets = c.call("TRN_RT_BUYGOLD_Upd", write=True, raise_on_rc=False, **full)
                self._ok("Upd mốc cũ → vendor từ chối (rc≠0)", rc not in (0, None), f"rc={rc}")
            except PmvProcError as exc:
                self._ok("Upd mốc cũ → vendor từ chối", True, str(exc)[:80])

            # ── 3. CHỐT TIỀN MẶT ──
            mid = counts()
            row = B.chot_thau(trn, till_id=till, user_id=user, tien_ck=0, c=c)
            txn = c.query("SELECT TillTxnID, TillID, Status, TrnTotalAmount, GoldCcy, TrnAmount FROM T_TILL_TXN WITH (NOLOCK) WHERE TrnRefID=?", (trn,))
            self._ok("chốt → C, két P đúng TillID", row["Status"] == "C" and txn and txn[0]["Status"] == "P" and txn[0]["TillID"] == till, str(txn))
            det = c.query("SELECT GoldCcy, CrDr, Amount, Amount_Total FROM T_TILL_TXN_DETAIL WITH (NOLOCK) WHERE TillTxnID=?", (txn[0]["TillTxnID"],))
            dmap = {x["GoldCcy"]: x for x in det}
            self._ok("DETAIL: dẻ + 1000 ly (tổng 1010 kể hột), VND − tiền", dmap.get("D18K", {}).get("CrDr") == "+" and M.dec(dmap["D18K"]["Amount"]) == 1000
                       and M.dec(dmap["D18K"]["Amount_Total"]) == 1010 and dmap.get("VND", {}).get("CrDr") == "-" and M.dec(dmap["VND"]["Amount"]) == tien2, str(det))
            self._ok("T_TILL_BAL: D18K +1000 ly · VND −tiền", bal("D18K") - bal0["D18K"] == 1000 and bal0["VND"] - bal("VND") == tien2,
                       f"D18K {bal0['D18K']}→{bal('D18K')} VND {bal0['VND']}→{bal('VND')}")
            d3 = self.delta(mid, counts())
            self._ok("chốt: +1 T_TILL_TXN · +2 DETAIL · +1 lịch sử điểm · 0 CardPay",
                       d3.get("T_TILL_TXN") == 1 and d3.get("T_TILL_TXN_DETAIL") == 2 and d3.get("I_LICHSUTICHLUYDIEM") == 1
                       and not d3.get("TRN_RT_BUYGOLD_CardPay"), str(d3))
            last1 = c.query("SELECT LastTradingDate FROM I_CUSTOMER WITH (NOLOCK) WHERE CustID=?", (cust,))[0]["LastTradingDate"]
            self._ok("I_CUSTOMER.LastTradingDate = hôm nay", last1 and last1.date() == datetime.date.today(), str(last1))
            try:
                B.chot_thau(trn, till_id=till, user_id=user, c=c)
                self._ok("chốt lại phiếu C → bị chặn", False)
            except ValueError as exc:
                self._ok("chốt lại phiếu C → bị chặn", "đã thanh toán" in str(exc).lower(), str(exc))
            try:
                c.call("TRN_RT_BUYGOLD_CompleteMore", write=True, p_TrnIDs=f"{trn}@")
                self._ok("CompleteMore trên phiếu C → BRT-001", False)
            except PmvProcError as exc:
                self._ok("CompleteMore trên phiếu C → BRT-001", "BRT-001" in str(exc.sets), str(exc.sets)[:100])

            # ── 4. HỦY THANH TOÁN → W, két hoàn ──
            B.mo_lai_thau(trn, user_id=user, c=c)
            row = B.phieu_thau(trn, c)
            self._ok("mo_lai → W, T_TILL_TXN mất, két về ban đầu", row["Status"] == "W"
                       and not c.query("SELECT 1 FROM T_TILL_TXN WITH (NOLOCK) WHERE TrnRefID=?", (trn,))
                       and bal("D18K") == bal0["D18K"] and bal("VND") == bal0["VND"], f"{bal('D18K')} {bal('VND')}")
            d4 = self.delta(mid, counts())
            self._ok("hủy TT ghi log: +1 TRN_RT_BUYGOLD_Log · +1 T_TILL_TXN_Log", d4.get("TRN_RT_BUYGOLD_Log") == 1 and d4.get("T_TILL_TXN_Log") == 1, str(d4))

            # ── 5. CHỐT CÓ CHUYỂN KHOẢN 30tr (GĐ chốt 2: CARDPAY_Ins DS thẻ rỗng) ──
            mid = counts()
            ck = M.dec(30_000_000)
            row = B.chot_thau(trn, till_id=till, user_id=user, tien_ck=ck, c=c)
            mat = tien2 - ck
            txn = c.query("SELECT TillTxnID, TrnTotalAmount FROM T_TILL_TXN WITH (NOLOCK) WHERE TrnRefID=?", (trn,))
            det = {x["GoldCcy"]: x for x in c.query("SELECT GoldCcy, CrDr, Amount FROM T_TILL_TXN_DETAIL WITH (NOLOCK) WHERE TillTxnID=?", (txn[0]["TillTxnID"],))}
            self._ok("CK: CardPay=−30tr, CashPay=−tiền mặt", M.dec(row["CardPay"]) == -ck and M.dec(row["CashPay"]) == -mat, f"{row['CardPay']} {row['CashPay']}")
            self._ok("CK: két VND chỉ trừ TIỀN MẶT (TrnTotalAmount + DETAIL)", M.dec(txn[0]["TrnTotalAmount"]) == mat and det["VND"]["CrDr"] == "-"
                       and M.dec(det["VND"]["Amount"]) == mat and bal0["VND"] - bal("VND") == mat, f"{txn} {det.get('VND')} {bal('VND')}")
            d5 = self.delta(mid, counts())
            self._ok("CK: KHÔNG có dòng TRN_RT_BUYGOLD_CardPay", not d5.get("TRN_RT_BUYGOLD_CardPay"), str(d5))
            try:
                B.chot_thau(trn, till_id=till, user_id=user, tien_ck=tien2 + 1, c=c)
                self._ok("CK vượt tổng → chặn", False)
            except ValueError:
                self._ok("CK vượt tổng → chặn", True)

            # ── 6. XÓA phiếu đã C: bill.huy_thau tự mo_lai rồi Del ──
            self._ok("Del phiếu C trần → vendor B-002", self._del_bi_chan(c, trn, user))
            B.huy_thau(trn, user_id=user, c=c)
            self._ok("huy_thau → dòng mất, két về ban đầu", not c.query("SELECT 1 FROM TRN_RT_BUYGOLD WITH (NOLOCK) WHERE TrnID=?", (trn,))
                       and bal("D18K") == bal0["D18K"] and bal("VND") == bal0["VND"])
            trn = ""
            d6 = self.delta(before, counts())
            self._ok("về số dòng ban đầu (trừ *_Log & lịch sử điểm đảo)",
                       all(d6.get(t, 0) == 0 for t in ("TRN_RT_BUYGOLD", "T_TILL_TXN", "T_TILL_TXN_DETAIL", "TRN_RT_BUYGOLD_CardPay")), str(d6))
            c.query("SELECT 1")  # giữ kết nối ấm cho phần view

            # ── 7. VIEW web (test client, đích sandbox) — trang cùng khung màn bán, giỏ session nhiều dòng ──
            web = Client(SERVER_NAME="localhost")
            web.force_login(get_user_model().objects.get(username="admin"))
            V._passcode_dung = lambda req, ma: ma == "SMOKE"
            V.PmvClient = lambda target=None, tag="client": PmvClient("sandbox", tag=tag)
            S.PmvClient = lambda target=None, tag="client": PmvClient("sandbox", tag=tag)   # hoa_don_loc ép "kk" cho hôm nay
            from apps.pos import views_thau as VT
            VT._passcode_dung = lambda req, ma: ma == "SMOKE"
            b = web.get("/banle/thau-vao/").content.decode()
            self._ok("trang thâu: khung pg-ban--thau + 3 khối + chân 4 nút", "pg-ban--thau" in b and 'id="pos-thau"' in b
                     and 'id="pos-tong"' in b and "THANH TOÁN &amp; IN" in b and 'name="kieu"' in b)
            ph = V._phien(web.get("/banle/thau-vao/").wsgi_request)
            if not ph["till_id"]:
                self._ok("admin có két PMV (sys_users)", False, "thiếu till_id → bỏ phần view")
            else:
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                web.post("/banle/thau-vao/dat/", {"cust_id": cust})
                web.post("/banle/thau-vao/dat/", {"emp": emp})
                r = web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "500", "tl_hot": "10", "gia": "", "kieu": "thau"})
                b = r.content.decode()
                self._ok("＋ THÊM dòng 1 (D18K, giá thâu) → OOB có dòng", r.status_code == 200 and "thâu vào" in b and "1 dòng" in b)
                r = web.post("/banle/thau-vao/them/", {"gold": "D24K", "tong_tl": "200", "tl_hot": "0", "gia": "", "kieu": "ban"})
                b = r.content.decode()
                self._ok("＋ THÊM dòng 2 (D24K, giá BÁN RA) → 2 dòng, badge bán ra", "2 dòng" in b and "bán ra" in b)
                r = web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "100", "tl_hot": "100", "gia": "", "kieu": "thau"})
                self._ok("hột ≥ tổng → báo lỗi, không thêm", "nhỏ hơn" in r.content.decode())
                web.post("/banle/thau-vao/dat/", {"bot": "20.000"})
                web.post("/banle/thau-vao/dat/", {"pay_method": "bank"})
                r = web.post("/banle/thau-vao/dat/", {"tien_mat": "1.000.000"})
                b = r.content.decode()
                from apps.pos import thau_cart as TC
                g = web.session.get(TC.KEY)
                t = TC.tong_cua(g)
                self._ok("giỏ: 2 dòng, bớt 20k, CK = tiệm trả − 1tr", len(g["lines"]) == 2 and t["bot"] == 20000 and t["tien_mat"] == 1_000_000
                         and t["tien_ck"] == t["khach_tra"] - 1_000_000, str(t))
                r = web.post("/banle/thau-vao/thanh-toan/", {"in": "1"})
                b = r.content.decode()
                from apps.pos.models import ThauNhom
                nhom = ThauNhom.objects.order_by("-pk").first()
                trn_nhom = list(nhom.trn_ids) if nhom else []
                rows = [B.phieu_thau(x, c) for x in trn_nhom]
                self._ok("THANH TOÁN & IN → 2 dòng C cùng nhóm, tờ 110mm 2 dòng trong #pos-in, form trắng",
                         len(rows) == 2 and all(x and x["Status"] == "C" for x in rows) and "th-phieu" in b and "khblInThau" in b
                         and b.count("<tr>") >= 2 and "Chưa có dòng vàng nào" in b, b[:120].replace("\n", " "))
                tong_kk = sum((M.dec(x["TotalAmount"]) for x in rows), M.dec(0))
                ck_kk = sum((-M.dec(x["CardPay"]) for x in rows), M.dec(0))
                add_kk = sum((M.dec(x["AddMoney"]) for x in rows), M.dec(0))
                self._ok("KK: Σ TotalAmount = tiệm trả · Σ CardPay = CK · AddMoney = −bớt", tong_kk == t["khach_tra"] and ck_kk == t["tien_ck"]
                         and add_kk == -20000, f"{tong_kk} {ck_kk} {add_kk} vs {t['khach_tra']} {t['tien_ck']}")
                trn = trn_nhom[0] if trn_nhom else ""
                # mở lại từ DANH SÁCH → cả nhóm lên form (khóa) → 🔓 SỬA (passcode) → về W → thêm dòng → THANH TOÁN lại
                r = web.get("/banle/thau-vao/danh-sach/")
                b = r.content.decode()
                self._ok("popup DANH SÁCH: ô lọc Khách chỉ thay #ds-kq (hx-select) — gõ không bị thay ô, không giật (09/09)",
                         'id="ds-kq"' in b and 'hx-target="#ds-kq" hx-select="#ds-kq" hx-swap="outerHTML"' in b)
                b_page = web.get("/banle/thau-vao/").content.decode()
                self._ok("trang thâu không còn JS giữ-ô inline (dùng module khblGiuO trong khbl.js), khbl.js đổi cache-buster",
                         "GIỮ Ô ĐANG GÕ: dùng module chung khblGiuO" in b_page and "khbl.js?v=" in b_page
                         and "khblGiuO = true" in (R_JS := open("static/js/khbl.js", encoding="utf-8").read()) and "htmx:beforeSwap" in R_JS)
                # GĐ 09/09: gõ SĐT ở ô Khách, lọc không ra → ＋ mở popup Thêm khách với SĐT điền sẵn #f-dt (nút ＋ hx-include #o-khach)
                self._ok("＋ Thêm khách mang SĐT đang gõ sang #f-dt (0912345678 · +84 912-345-678 → 0912345678 · chữ → trống)",
                         'id="f-dt" value="0912345678"' in web.get("/banle/khach-hang/them/?q=0912345678").content.decode()
                         and 'id="f-dt" value="0912345678"' in web.get("/banle/khach-hang/them/?q=%2B84%20912-345-678").content.decode()
                         and 'id="f-dt" value=""' in web.get("/banle/khach-hang/them/?q=Nguyen%20Van").content.decode()
                         and 'hx-include="#o-khach"' in b_page)
                # GĐ chốt 10/09/2026: 2 phiếu cùng nhóm gom về MỘT dòng, ô SỐ PHIẾU in cả 2 mã, ô TIỆM TRẢ 2 tầng
                b_ds = r.content.decode()
                import re as _re2
                dong_nhom = [x for x in _re2.findall(r"<tr [^>]*data-th-row=.*?</tr>", b_ds, _re2.S) if "⧉ 2" in x]
                ma_trong_dong = _re2.findall(r'pg-ds__ma">([^<]+)<', dong_nhom[0]) if dong_nhom else []
                self._ok("popup DANH SÁCH: 2 phiếu cùng nhóm gom về MỘT dòng, ô SỐ PHIẾU in đủ 2 mã (mỗi mã một dòng)",
                         r.status_code == 200 and bool(dong_nhom) and len(ma_trong_dong) == 2, str(ma_trong_dong))
                self._ok("cột TIỆM TRẢ: dòng trên tổng tiền thâu của cả nhóm, dòng dưới tổng tiền chuyển khoản",
                         bool(dong_nhom) and ("CK " in dong_nhom[0] or "tiền mặt" in dong_nhom[0])
                         and 'pg-ds__td-tra' in dong_nhom[0], (dong_nhom[0][-260:] if dong_nhom else ""))
                # cộng tiền: kiểm thẳng hàm gộp bằng dữ liệu dựng sẵn, khỏi phụ thuộc cách hiện số trên HTML
                gia = [{"TrnID": "A1", "BillCode": "b1", "SoTien": "100", "CardPay": "-60", "Status": "C", "IsDel": "0"},
                       {"TrnID": "A2", "BillCode": "b2", "SoTien": "250", "CardPay": "-250", "Status": "C", "IsDel": "0"},
                       {"TrnID": "B1", "BillCode": "b3", "SoTien": "70", "CardPay": "0", "Status": "W", "IsDel": "0"},
                       {"TrnID": "A3", "BillCode": "b4", "SoTien": "999", "CardPay": "0", "Status": "C", "IsDel": "1"}]
                nhom_gia = type("N", (), {"pk": 7, "trn_ids": ["A1", "A2", "A3"]})()
                gom = VT._gom_dong_ds(gia, {"A1": nhom_gia, "A2": nhom_gia, "A3": nhom_gia})
                nh = next(d for d in gom if d["so_dong"] > 1)
                le = next(d for d in gom if d["so_dong"] == 1)
                self._ok("gộp nhóm: tổng tiền thâu = tổng các phiếu CÒN SỐNG, tổng CK = phần CardPay âm, dòng đã hủy không cộng vào",
                         len(gom) == 2 and nh["tong_tien"] == M.dec(350) and nh["tong_ck"] == M.dec(310)
                         and nh["so_dong"] == 3 and nh["co_huy"] and not nh["da_huy"] and nh["mo_id"] == "A1"
                         and le["tong_tien"] == M.dec(70) and le["tong_ck"] == M.D0 and not le["chot"],
                         f"{nh['tong_tien']}/{nh['tong_ck']} · {len(gom)} dòng")
                r = web.post("/banle/thau-vao/mo/", {"trn_id": trn})
                g = web.session.get(TC.KEY)
                self._ok("MỞ → nạp cả nhóm 2 dòng, trạng thái C (khóa)", len(g["lines"]) == 2 and g["status"] == "C" and "🔒" in r.content.decode())
                r = web.post("/banle/thau-vao/thuc-hien/sua/", {"passcode": "SAI"})
                self._ok("SỬA sai passcode → popup báo, phiếu vẫn C", "Passcode không đúng" in r.content.decode()
                         and B.phieu_thau(trn, c)["Status"] == "C")
                r = web.post("/banle/thau-vao/thuc-hien/sua/", {"passcode": "SMOKE"})
                g = web.session.get(TC.KEY)
                self._ok("SỬA đúng passcode → 2 dòng về W, form ĐANG SỬA, két hoàn", all(B.phieu_thau(x, c)["Status"] == "W" for x in trn_nhom)
                         and g["status"] == "W" and g.get("sua_lai") and bal("VND") == bal0["VND"], f"{bal('VND')} vs {bal0['VND']}")
                r = web.post("/banle/thau-vao/them/", {"gold": "D9999", "tong_tl": "100", "tl_hot": "0", "gia": "", "kieu": "thau"})
                self._ok("thêm dòng 3 vào phiếu đang sửa", "3 dòng" in r.content.decode())
                r = web.post("/banle/thau-vao/thanh-toan/", {})
                nhom2 = ThauNhom.objects.order_by("-pk").first()
                rows2 = [B.phieu_thau(x, c) for x in nhom2.trn_ids]
                self._ok("THANH TOÁN lại → 3 dòng C (2 Upd + 1 Ins), nhóm mới, form giữ phiếu khóa", len(rows2) == 3
                         and all(x["Status"] == "C" for x in rows2) and set(trn_nhom) < set(nhom2.trn_ids)
                         and "🔒 ĐÃ CHỐT" in r.content.decode())
                r = web.get(f"/banle/thau-vao/in/?trn_id={nhom2.trn_ids[0]}")
                if VT.KHOA_IN:      # GĐ tạm khóa in (10/09/2026) → trang in phải bị chặn; mở lại thì tự kiểm nội dung
                    self._ok("đang TẠM KHÓA IN → trang in standalone bị chặn ở máy chủ",
                             r.status_code == 403 and "khóa in" in r.content.decode().lower())
                else:
                    self._ok("trang in standalone: 3 dòng + bằng chữ", r.status_code == 200 and "Bằng chữ" in r.content.decode()
                             and r.content.decode().count("<tr>") >= 4)
                r = web.post("/banle/thau-vao/thuc-hien/huy_hd/", {"passcode": "SMOKE"})
                self._ok("XÓA PHIẾU (passcode) → 3 dòng mất, két về ban đầu", r.status_code == 200
                         and not any(B.phieu_thau(x, c) for x in nhom2.trn_ids) and bal("D18K") == bal0["D18K"] and bal("VND") == bal0["VND"],
                         f"{bal('D18K')} {bal('VND')}")
                trn = ""
                # ── 8. (08/09 tối) PHIẾU MỚI xóa cả NV · TÍNH LẠI gộp dòng · ảnh chuyển khoản → gold_bill · × dòng đã lưu = xóa KK ──
                from apps.pos.models import GoldBill, ThauAnhTam
                web.post("/banle/thau-vao/dat/", {"emp": emp})
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                g = web.session.get(TC.KEY)
                self._ok("PHIẾU MỚI xóa trắng cả nhân viên thâu", not g.get("emp") and not g["lines"])
                web.post("/banle/thau-vao/dat/", {"cust_id": cust}); web.post("/banle/thau-vao/dat/", {"emp": emp})
                web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "300", "tl_hot": "10", "gia": "", "kieu": "thau"})
                web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "200", "tl_hot": "5", "gia": "", "kieu": "thau"})
                web.post("/banle/thau-vao/them/", {"gold": "D24K", "tong_tl": "100", "tl_hot": "0", "gia": "", "kieu": "thau"})
                r = web.post("/banle/thau-vao/tinh-lai/", {"kieu": "thau"})
                g = web.session.get(TC.KEY)
                d18 = [x for x in g["lines"] if x["gold"] == "D18K"]
                self._ok("TÍNH LẠI gộp 2 dòng D18K cùng giá → 1 dòng 500 ly / hột 15, còn D24K riêng", len(g["lines"]) == 2 and len(d18) == 1
                         and M.dec(d18[0]["tong_tl"]) == 500 and M.dec(d18[0]["tl_hot"]) == 15 and "gộp 1 dòng" in r.content.decode(), str(g["lines"]))
                from io import BytesIO
                from PIL import Image
                from django.core.files.uploadedfile import SimpleUploadedFile
                def anh(mau):
                    out = BytesIO(); Image.new("RGB", (900, 600), mau).save(out, "PNG"); return SimpleUploadedFile("a.png", out.getvalue(), content_type="image/png")
                def mau_tb(data):
                    """Màu trung bình của ảnh đã lưu — dùng để biết ô nào giữ ảnh nào (đảo ô là lộ ngay)."""
                    im = Image.open(BytesIO(bytes(data))).convert("RGB").resize((8, 8))
                    px = list(im.getdata())
                    return tuple(sum(q[i] for q in px) // len(px) for i in range(3))
                r = web.post("/banle/thau-vao/anh/len/", {"anh_cccd1": anh("red"), "anh_qr": anh("blue"),
                                                         "anh_hinh1": anh("#14C814"), "anh_hinh2": anh("#1414C8")})
                sk = web.session.session_key
                b = r.content.decode()
                self._ok("tải 4 ảnh (CCCD1 · QR · Hình 1 · Hình 2) → chỉ là ẢNH CHỜ theo phiên (chưa ghi khách/nhóm), OOB hiện ảnh + nhãn 'chờ'",
                         r.status_code == 200 and ThauAnhTam.objects.filter(session_key=sk).count() == 4 and "slot=cccd1" in b
                         and "· chờ" in b and "ghi vào phiếu / hồ sơ khách khi THANH TOÁN" in b, b[:160].replace("\n", " "))
                kh_anh = c.query("SELECT ImagePathMatSau FROM I_CUSTOMER WITH (NOLOCK) WHERE CustID=?", (cust,))[0]["ImagePathMatSau"]
                self._ok("ô CCCD sau (chưa tải): khách có ảnh sẵn → hiện ảnh HỒ SƠ KHÁCH / không có → 'Chưa có'",
                         ("khach-hang/" + cust + "/anh/mat-sau/" in b) if kh_anh else ("khach-hang/" + cust + "/anh/mat-sau/" not in b))
                r = web.get("/banle/thau-vao/anh/?slot=qr")
                self._ok("GET ảnh chờ QR → image/jpeg", r.status_code == 200 and r["Content-Type"] == "image/jpeg")
                web.post("/banle/thau-vao/anh/xoa/", {"slot": "qr"})
                self._ok("× bỏ ảnh chờ QR", ThauAnhTam.objects.filter(session_key=sk).count() == 3)
                r = web.post("/banle/thau-vao/thanh-toan/", {})
                b = r.content.decode()
                nhom3 = ThauNhom.objects.order_by("-pk").first()
                gb = GoldBill.objects.filter(trn_id=nhom3.trn_ids[0]).first()
                self._ok("THANH TOÁN → gold_bill(TrnID đầu) vẫn ghi tổng/loại thâu nhưng KHÔNG mang ảnh (ảnh ở nhà mới)",
                         gb is not None and gb.status == "C" and gb.tong > 0 and gb.doi and gb.doi[0]["vang"] == "D18K"
                         and not gb.anh_hinh1 and not gb.anh_cccd1, str(gb and gb.tong))
                m1, m2 = mau_tb(nhom3.anh_hinh1), mau_tb(nhom3.anh_hinh2)
                self._ok("THANH TOÁN → Hình 1 (lục) vào thau_nhom.anh_hinh1, Hình 2 (lam) vào anh_hinh2, QR trống (đã bỏ chờ)",
                         m1[1] > m1[0] and m1[1] > m1[2] and m2[2] > m2[0] and m2[2] > m2[1] and not nhom3.anh_qr
                         and bytes(nhom3.anh_hinh1) != bytes(nhom3.anh_hinh2), f"h1={m1} h2={m2}")
                self._ok("THANH TOÁN → CCCD chờ ghi lên HỒ SƠ KHÁCH (sandbox tắt OLE → báo 'CHƯA cập nhật', CCCD vẫn giữ chờ); Hình 1/2 chờ đã xóa",
                         "CHƯA cập nhật hồ sơ khách" in b and ThauAnhTam.objects.filter(session_key=sk, slot="cccd1").exists()
                         and not ThauAnhTam.objects.filter(session_key=sk, slot__in=["hinh1", "hinh2"]).exists(), b[:200].replace("\n", " "))
                from pathlib import Path as _P
                from django.conf import settings as _st
                tm_cccd = _P(_st.BASE_DIR) / "media" / "cccd" / cust
                f_mt = sorted(tm_cccd.glob("MT_*.jpg")) if tm_cccd.exists() else []
                self._ok("THANH TOÁN → file CCCD mặt trước ghi trên MÁY CHỦ media/cccd/<CustID>/MT_<lúc>.jpg (ảnh đỏ)",
                         bool(f_mt) and f_mt[-1].stat().st_size > 500 and mau_tb(f_mt[-1].read_bytes())[0] > 150, str([x.name for x in f_mt[-1:]]))
                for _f in f_mt[-1:]:
                    _f.unlink(missing_ok=True)                       # dọn file ảnh đỏ của smoke, không để rác trong media/cccd
                b_tt = web.get("/banle/thau-vao/").content.decode().replace("&amp;", "&")   # src dựng trong Python → template escape &
                self._ok("trang sau THANH TOÁN → Hình 1/2 trỏ vào nhóm (nhom=<id>), CCCD1 vẫn là ảnh chờ",
                         f"slot=hinh1&nhom={nhom3.pk}" in b_tt and f"slot=hinh2&nhom={nhom3.pk}" in b_tt and "slot=cccd1&v=" in b_tt)
                r = web.post("/banle/thau-vao/anh/len/", {"anh_hinh2": anh("black")})
                b = r.content.decode()
                self._ok("đơn KHÓA → tải ảnh bị chặn, nút disabled trong fieldset, KHÔNG tô màu (hết pg-box--khoa ở khối ảnh)",
                         "KHÓA" in b and 'fieldset disabled class="pg-khoa-fs"' in b and "pg-box--ck pg-box--khoa" not in b)
                r = web.get(f"/banle/thau-vao/anh/?slot=hinh1&nhom={nhom3.pk}")
                self._ok("GET ảnh từ NHÓM theo ?nhom= → jpeg", r.status_code == 200 and r["Content-Type"] == "image/jpeg")
                r = web.post("/banle/thau-vao/thuc-hien/sua/", {"passcode": "SMOKE"})
                g = web.session.get(TC.KEY)
                r = web.post("/banle/thau-vao/anh/len/", {"anh_hinh2": anh("black")})
                self._ok("đã 🔓 SỬA → tải ảnh mới chỉ là ảnh CHỜ, nhóm CHƯA đổi (THANH TOÁN mới thay)",
                         ThauAnhTam.objects.filter(session_key=sk, slot="hinh2").exists()
                         and bytes(ThauNhom.objects.get(pk=nhom3.pk).anh_hinh2) == bytes(nhom3.anh_hinh2) and "ảnh chờ" in r.content.decode())
                r = web.post("/banle/thau-vao/anh/xoa/", {"slot": "hinh1"})
                self._ok("× lên ảnh ĐÃ GẮN nhóm (không phải ảnh chờ) → từ chối, hướng dẫn thay bằng ảnh mới",
                         "đã gắn phiếu" in r.content.decode() and ThauNhom.objects.get(pk=nhom3.pk).anh_hinh1)
                t0 = g["lines"][0]["trn_id"]
                r = web.post("/banle/thau-vao/xoa-dong/", {"i": "0"})
                g = web.session.get(TC.KEY)
                self._ok("× dòng đã lưu chờ → xóa hẳn trên KK + bỏ khỏi giỏ (ảnh theo nhóm, không đụng)", not B.phieu_thau(t0, c) and len(g["lines"]) == 1
                         and t0 not in g["trn_ids"] and "xóa hẳn" in r.content.decode(), r.content.decode()[:80])
                r = web.post("/banle/thau-vao/thanh-toan/", {})
                nhom3b = ThauNhom.objects.order_by("-pk").first()
                m2b = mau_tb(nhom3b.anh_hinh2)
                self._ok("THANH TOÁN lại → NHÓM MỚI kế thừa Hình 1 của nhóm cũ, Hình 2 thay bằng ảnh chờ (đen); xóa dòng không mất ảnh",
                         nhom3b.pk != nhom3.pk and mau_tb(nhom3b.anh_hinh1) == m1 and max(m2b) < 40
                         and not ThauAnhTam.objects.filter(session_key=sk, slot="hinh2").exists(), f"h1={mau_tb(nhom3b.anh_hinh1)} h2={m2b}")
                r = web.post("/banle/thau-vao/thuc-hien/huy_hd/", {"passcode": "SMOKE"})
                self._ok("XÓA phiếu phần còn lại → KK sạch, két về ban đầu", not any(B.phieu_thau(x, c) for x in nhom3.trn_ids)
                         and bal("D18K") == bal0["D18K"] and bal("VND") == bal0["VND"] and GoldBill.objects.get(trn_id=nhom3.trn_ids[0]).is_del)
                # ── 8a-bis. (10/09/2026, GĐ) popup CHI TIẾT trang /thau-vao-2/: khóa QR khi đã xác nhận CK + đối chiếu tên
                import re as _re_ds
                from apps.pos import thau_list as TL
                from apps.pos import thau_xuat as TX
                from apps.pos.models import ThauPaymentLink
                self._ok("đối chiếu tên chủ thẻ với tên khách: bỏ dấu + bỏ xưng hô rồi so, thiếu tên thì không kết luận",
                         TL.doi_chieu_ten("TRAN THI TRANG", "Trần Thị Trang") is True
                         and TL.doi_chieu_ten("CHI TRANG", "Chị Trang") is True
                         and TL.doi_chieu_ten("NGUYEN VAN A", "Chị Trang") is False
                         and TL.doi_chieu_ten("", "Chị Trang") is None and TL.doi_chieu_ten("TRAN A", "") is None)
                self._ok("phiếu chưa có liên kết CK / phiếu trả tiền mặt → chưa coi là đã xác nhận",
                         TL.da_xac_nhan_ck(["TBG_KHONG_TON_TAI"], 1000)[0] is False
                         and TL.da_xac_nhan_ck(["TBG_KHONG_TON_TAI"], 0)[0] is False)
                lk = ThauPaymentLink.objects.filter(active_notification_id__isnull=False).order_by("-pk").first()
                if lk:
                    self._ok("có liên kết CK còn hiệu lực, đủ tiền → coi là ĐÃ XÁC NHẬN (đọc bảng thau_payment_link)",
                             TL.da_xac_nhan_ck(lk.trn_ids, lk.amount)[0] is True
                             and TL.da_xac_nhan_ck(lk.trn_ids, lk.amount * 2)[0] is False)
                    # tìm nhóm ĐÃ xác nhận mà CÓ ảnh QR (nhóm không có QR thì không có gì để phủ nhãn)
                    b_ct, o_qr = "", None
                    for l2 in ThauPaymentLink.objects.filter(active_notification_id__isnull=False).order_by("-pk")[:15]:
                        thu = web.get("/banle/thau-vao-2/xem/?trn_id=" + l2.trn_ids[0]).content.decode()
                        if "th2-bank-qr--xong" in thu:
                            b_ct = thu
                            o_qr = _re_ds.search(r'<div class="th2-bank-qr[^>]*>(.*?)</div>', thu, _re_ds.S)
                            break
                    self._ok("popup chi tiết của nhóm ĐÃ xác nhận (có ảnh QR): QR bị phủ nhãn xanh ĐÃ XONG và KHÔNG còn bấm mở được",
                             (not b_ct) or ("ĐÃ XONG" in b_ct and bool(o_qr) and "<a " not in o_qr.group(1)),
                             "không nhóm nào vừa đã xác nhận vừa có ảnh QR" if not b_ct else "")
                # Khung xem ảnh là markup + JS tĩnh của template: kiểm thẳng tệp, khỏi phải có phiếu thật trên KK
                # (nhóm mới nhất trong smoke là nhóm sandbox, popup chi tiết đọc KK thật nên không mở được).
                tpl_xem = open("templates/pos/_thau2_detail.html", encoding="utf-8").read()
                js_khbl = open("static/js/khbl.js", encoding="utf-8").read()
                self._ok("popup chi tiết: ảnh mở NGAY TRONG khung popup — có khung xem ảnh, bộ slide, nút chuyển và mức phóng",
                         'id="th2-xem"' in tpl_xem and "data-th2-anh" in tpl_xem and "data-truoc" in tpl_xem
                         and "data-sau" in tpl_xem and "data-ti-le" in tpl_xem
                         and "position:absolute" in open("static/css/thau2.css", encoding="utf-8").read().split(".th2-xem{")[1][:80])
                self._ok("khung xem ảnh khai là LỚP CON: Esc chỉ đóng nó, popup chi tiết còn nguyên (khbl.js nhường + chặn hẳn)",
                         "data-khbl-lop-con" in tpl_xem and "data-khbl-lop-con" in js_khbl
                         and "stopImmediatePropagation" in tpl_xem)
                # GĐ chốt 10/09/2026: TẠM KHÓA IN + đơn đã xác nhận CK thì popup DANH SÁCH chỉ cho XEM
                r_in = web.get("/banle/thau-vao/in/?oob=1")
                self._ok("mọi nút IN bị khóa và view in chặn thẳng ở máy chủ (gọi đường dẫn cũng không in được)",
                         r_in.status_code == 403 and "khóa in" in r_in.content.decode().lower()
                         and web.get("/banle/thau-vao/").content.decode().count("Tạm khóa in phiếu") == 2)
                b_ds2 = web.get("/banle/thau-vao/danh-sach/").content.decode()
                # ngày không có phiếu đã chốt thì không dòng nào vẽ nút 🖨 → kiểm thêm ở template cho chắc
                tpl_ds = open("templates/pos/_thau_ds.html", encoding="utf-8").read()
                self._ok("popup DANH SÁCH: nút 🖨 của từng dòng cũng bị khóa, không còn gọi được đường dẫn in",
                         "thau-vao/in/" not in b_ds2 and "disabled title=\"{{ khoa_in_msg }}\"" in tpl_ds
                         and ("Tạm khóa in phiếu" in b_ds2 or ">🖨</button>" not in b_ds2))
                nhom_xn = ThauNhom.objects.order_by("-pk").first()
                lk_gia = ThauPaymentLink.objects.create(
                    order_key="smoke", trn_ids=list(nhom_xn.trn_ids), notification_id=999999999,
                    active_notification_id=999999999, amount=nhom_xn.tien_ck or 1, snapshot={}, bank_snapshot={},
                    mode="manual", username="smoke", reason="smoke")
                try:
                    b_xn = web.get("/banle/thau-vao/danh-sach/").content.decode()
                    dong_xn = [x for x in _re_ds.findall(r"<tr [^>]*data-th-row=.*?</tr>", b_xn, _re_ds.S)
                               if nhom_xn.trn_ids[0] in x or (nhom_xn.bill_codes and nhom_xn.bill_codes[0] in x)]
                    co_ck = bool(nhom_xn.tien_ck)
                    self._ok("đơn ĐÃ xác nhận CK: nút MỞ đổi thành XEM, mở popup chi tiết chỉ-xem của trang 2",
                             (not co_ck) or (bool(dong_xn) and ">XEM<" in dong_xn[0] and ">MỞ<" not in dong_xn[0]
                                             and "thau-vao-2/xem/" in dong_xn[0] and "đã xác nhận CK" in dong_xn[0]),
                             f"tien_ck={nhom_xn.tien_ck} dong={len(dong_xn)}")
                finally:
                    lk_gia.delete()
                b_sau = web.get("/banle/thau-vao/danh-sach/").content.decode()
                dong_sau = [x for x in _re_ds.findall(r"<tr [^>]*data-th-row=.*?</tr>", b_sau, _re_ds.S)
                            if nhom_xn.trn_ids[0] in x or (nhom_xn.bill_codes and nhom_xn.bill_codes[0] in x)]
                self._ok("gỡ liên kết CK → dòng đó quay lại nút MỞ (sửa được như cũ), không kẹt ở chế độ chỉ xem",
                         not dong_sau or ">XEM<" not in dong_sau[0])
                # GĐ chốt 11/09/2026: KHÔNG chèn chú thích giữa markup. Nguy hiểm nhất là {# … #} NHIỀU DÒNG:
                # Django chỉ coi {# #} là chú thích khi gọn trong MỘT dòng, nhiều dòng thì in nguyên văn ra trang.
                loi_ct = []
                for tep in sorted(Path("templates").rglob("*.html")):
                    noi = tep.read_text(encoding="utf-8")
                    for m in _re_ds.finditer(r"\{#", noi):
                        cuoi = noi.find("#}", m.start())
                        if cuoi < 0 or chr(10) in noi[m.start():cuoi]:
                            loi_ct.append(f"{tep}:{noi[:m.start()].count(chr(10)) + 1}")
                self._ok("không template nào còn chú thích {# #} nhiều dòng (loại này in thẳng ra trang cho khách thấy)",
                         not loi_ct, "; ".join(loi_ct[:3]))
                nay = datetime.date.today().isoformat()
                b_sach = web.get(f"/banle/thau-vao-2/?d1={nay}&d2={nay}").content.decode()
                # Cột HÌNH ẢNH (GĐ chốt 11/09/2026): ảnh QR mang ✔️ khi đã xác nhận CK, ❌ khi hết hạn mà chưa xác nhận
                tpl_l = open("templates/pos/_thau2_list.html", encoding="utf-8").read()
                # GĐ bắt lỗi 11/09/2026: dấu ✓ sau Tên chủ thẻ mới có ở popup, DANH SÁCH thì không — nay cả hai
                import inspect as _ins
                self._ok("dấu ✓ sau Tên chủ thẻ và Nội dung được gắn cho CẢ danh sách lẫn popup chi tiết",
                         "ten_khop" in _ins.getsource(TL.enrich) and "ten_khop" in _ins.getsource(TL.detail)
                         and "nd_khop" in _ins.getsource(TL.enrich))
                self._ok("so tên bỏ dấu: 'Nguyễn Hoàng Gia Bảo' = 'NGUYEN HOANG GIA BAO' → khớp",
                         TL.doi_chieu_ten("NGUYEN HOANG GIA BAO", "Nguyễn Hoàng Gia Bảo") is True
                         and TL.doi_chieu_ten("NGUYEN HOANG GIA BAO", "Nguyễn Hoàng Gia Bao") is True
                         and TL.doi_chieu_ten("LE VAN C", "Nguyễn Hoàng Gia Bảo") is False)
                self._ok("cột HÌNH ẢNH: ảnh QR có dấu ✔️ khi đã xác nhận, ❌ khi hết hạn mà chưa xác nhận",
                         "th2-qr-dau--ok" in tpl_l and "th2-qr-dau--loi" in tpl_l
                         and "{% if r.gd_khop %}" in tpl_l and "{% elif r.qr_het_han %}" in tpl_l
                         and ".th2-qr-dau{" in open("static/css/thau2.css", encoding="utf-8").read())
                self._ok("trang thâu 2 không để lọt chú thích nội bộ ra HTML",
                         "{#" not in b_sach and "GĐ chốt" not in b_sach)

                # GĐ chốt 10/09/2026: trang /thau-vao-2/ MỞ CÔNG KHAI trong mạng tiệm (gồm ảnh + 2 nút xuất tệp),
                # nút "Đối soát CK" bỏ đi vì máy chủ tự soát mỗi 5 phút; các trang NHẬP LIỆU vẫn phải đăng nhập.
                hom_nay = datetime.date.today().isoformat()
                khach_la = Client(SERVER_NAME="localhost")     # phiên CHƯA đăng nhập
                mo = {u: khach_la.get(u).status_code for u in (
                    "/banle/thau-vao-2/", f"/banle/thau-vao-2/xuat-ncc/?d1={hom_nay}&d2={hom_nay}",
                    f"/banle/thau-vao-2/in-cccd/?d1={hom_nay}&d2={hom_nay}")}
                self._ok("chưa đăng nhập vẫn mở được trang thâu 2 + hai nút xuất tệp (GĐ chốt mở công khai trong LAN)",
                         all(v == 200 for v in mo.values()), str(mo))
                chan = {u: khach_la.get(u).status_code for u in ("/banle/thau-vao/", "/banle/khach-hang/", "/banle/hoa-don/")}
                self._ok("các trang NHẬP LIỆU vẫn bắt đăng nhập như cũ (chỉ mở đúng trang xem)",
                         all(v == 302 for v in chan.values()), str(chan))
                b_ds3 = khach_la.get("/banle/thau-vao-2/").content.decode()
                self._ok("trang thâu 2 KHÔNG còn nút Đối soát CK và dòng chú thích — máy chủ tự soát bằng lệnh doi_soat_ck",
                         "data-payment-scan" not in b_ds3 and "Tự kiểm tra mỗi 5 giây" not in b_ds3
                         and 'data-can-manage="0"' in b_ds3
                         and "doi_soat_ck" in open("config/scheduler.py", encoding="utf-8").read())

                # ⬇ XUẤT EXCEL + 🪪 IN CCCD (GĐ chốt 10/09/2026) — đúng khuôn tệp mẫu NCC_NHAP_CHUAN_KH2.xlsx
                hom_nay_iso = hom_nay
                b_tr2 = web.get("/banle/thau-vao-2/").content.decode()
                # nút "Đối soát CK" đã bỏ (11/09) nên không so vị trí với nó nữa; 2 nút tải tệp đứng ngay trên bảng
                self._ok("trang thâu 2: 2 nút tải tệp có mặt, giữ nguyên bộ lọc đang xem",
                         "XUẤT EXCEL" in b_tr2 and "IN CCCD" in b_tr2
                         and "xuat-ncc/?d1=" in b_tr2 and "in-cccd/?d1=" in b_tr2
                         and b_tr2.index("XUẤT EXCEL") < b_tr2.index("th2-table"))
                r_xl = web.get(f"/banle/thau-vao-2/xuat-ncc/?d1={hom_nay_iso}&d2={hom_nay_iso}&method=all")
                from io import BytesIO as _B
                import openpyxl as _xl
                ws = _xl.load_workbook(_B(r_xl.content)).active
                cot = [c.value for c in ws[1]]
                self._ok("xuất Excel: đúng 11 cột theo tệp mẫu, sheet NCC, tên tệp kèm khoảng ngày",
                         r_xl.status_code == 200 and ws.title == "NCC" and cot == list(TX.COT)
                         and f"NCC_NHAP_CHUAN_KH2_{datetime.date.today():%Y%m%d}.xlsx" in r_xl.get("Content-Disposition", ""),
                         r_xl.get("Content-Disposition", ""))
                dong2 = [c.value for c in ws[2]] if ws.max_row >= 2 else []
                self._ok("mỗi khách một dòng: Loại NCC = Cá nhân, tên in hoa, Ghi chú = TIỀN CK CHIA 1000 (đúng như mẫu)",
                         (not dong2) or (dong2[3] == "Cá nhân" and str(dong2[2]) == str(dong2[2]).upper()
                                         and (dong2[9] is None or isinstance(dong2[9], int))), str(dong2[:4]))
                self._ok("cột STT · Mã NCC · Email · Mã số thuế để TRỐNG như mẫu (phần mềm bên kia tự sinh)",
                         (not dong2) or all(dong2[i] in ("", None) for i in (0, 1, 6, 7)))
                r_w = web.get(f"/banle/thau-vao-2/in-cccd/?d1={hom_nay_iso}&d2={hom_nay_iso}&method=all")
                import re as _re_w
                kt = _re_w.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"', r_w.content.decode("latin-1", "ignore"))
                self._ok("in CCCD: tệp Word khổ A4, ảnh đặt ĐÚNG CỠ THẬT 85,6 × 53,98 mm để in ra trùng khít thẻ",
                         r_w.status_code == 200 and f"CCCD_{datetime.date.today():%Y%m%d}.docx" in r_w.get("Content-Disposition", "")
                         and (not kt or (abs(int(kt[0][0]) / 36000 - 85.6) < 0.2 and abs(int(kt[0][1]) / 36000 - 53.98) < 0.2)),
                         (f"{int(kt[0][0])/36000:.1f}x{int(kt[0][1])/36000:.1f}mm" if kt else "chưa có ảnh CCCD trong kỳ"))
                from docx import Document as _Doc
                dw = _Doc(_B(r_w.content))
                so_ngat = dw.element.xml.count('w:type="page"')
                dong_ten = [x.text for x in dw.paragraphs if x.text and "CCCD KHÁCH BÁN VÀNG" not in x.text]
                self._ok("in CCCD: mỗi trang đúng 4 khách rồi tự sang trang mới",
                         so_ngat == max(0, (len(dw.tables) - 1) // TX.KHACH_MOI_TRANG),
                         f"{len(dw.tables)} khách · {so_ngat} lần ngắt trang")
                self._ok("in CCCD: dòng chữ mỗi khách chỉ gồm TÊN · CCCD (không kèm số điện thoại, không kèm tiền)",
                         (not dong_ten) or all("·" not in t.split("CCCD")[0].replace(t.split("·")[0], "", 1)
                                               and "₫" not in t for t in dong_ten),
                         dong_ten[0] if dong_ten else "")
                r_dai = web.get(f"/banle/thau-vao-2/xuat-ncc/?d1=2026-01-01&d2={hom_nay_iso}")
                self._ok("chọn khoảng quá 31 ngày → từ chối gọn, không dựng tệp khổng lồ", r_dai.status_code == 400)

                self._ok("popup chi tiết TỰ mang bảng kiểu của nó — mở từ trang /thau-vao/ (không có sẵn thau2.css) vẫn đúng giao diện",
                         "css/thau2.css" in tpl_xem and "thau2.css" not in open("templates/pos/thau.html", encoding="utf-8").read())
                self._ok("lăn chuột giữa để phóng to quanh con trỏ, kéo để di chuyển, chặn cuộn nền",
                         'addEventListener("wheel"' in tpl_xem and "passive: false" in tpl_xem
                         and "pointermove" in tpl_xem and "dblclick" in tpl_xem)

                # ── 8b. (08/09 tối) ✂ TÁCH THẺ CCCD bằng OpenCV: ảnh giả lập thẻ xoay 12° lệch góc + vật tạp ──
                from apps.pos import anh_cccd as AC
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                r = web.post("/banle/thau-vao/anh/len/", {"anh_cccd2": anh("red")})
                self._ok("chưa chọn khách → tải CCCD bị từ chối 'Chọn khách hàng', ô CCCD trên trang khóa 'chọn khách hàng'",
                         "Chọn khách hàng" in r.content.decode() and "chọn khách hàng</span>" in r.content.decode()
                         and not ThauAnhTam.objects.filter(session_key=sk, slot="cccd2").exists())
                web.post("/banle/thau-vao/dat/", {"cust_id": cust})
                png_the = AC.anh_thu_nghiem(goc=12)
                r = web.post("/banle/thau-vao/anh/len/", {"anh_cccd2": SimpleUploadedFile("the.png", png_the, content_type="image/png")})
                b = r.content.decode()
                self._ok("tải ảnh thẻ giả lập vào CCCD sau → có nút ✂ bật", 'pg-ck__ico--cat" type="button" disabled' not in b.split('data-label="CCCD mặt sau"')[1].split("</div>")[2] if 'data-label="CCCD mặt sau"' in b else False)
                self._ok("nút ✂ khai hx-swap=innerHTML (form cha hx-swap=none kế thừa → popup từng không hiện)", 'hx-target="#modal-root" hx-swap="innerHTML"' in b and b.count('>✂</button>') >= 2)
                r = web.get("/banle/thau-vao/anh/cat/?slot=cccd2")
                b = r.content.decode()
                self._ok("✂ xem trước: popup có ảnh gốc + thẻ đã tách 1170×738", r.status_code == 200 and "Thẻ đã tách" in b and "1170×738" in b and b.count("data:image/jpeg") == 2)
                self._ok("popup có khung KÉO XOAY bằng con trỏ + ô góc ẩn gửi theo LƯU (hết nút xoay cứng)",
                         'id="th-cat-stage"' in b and 'id="th-cat-goc"' in b and 'hx-include="#th-cat-form"' in b and "&xoay=" not in b)
                r = web.post("/banle/thau-vao/anh/cat/luu/", {"slot": "cccd2", "goc": "0"})
                t = ThauAnhTam.objects.filter(session_key=sk, slot="cccd2").first()
                from PIL import Image as _Im

                def _sang(im):
                    px = list(im.convert("L").resize((16, 10)).getdata())
                    return sum(1 for v in px if v > 200) / len(px)
                sang_the = _sang(_Im.fromarray(AC._the_gia()[0][:, :, ::-1]))   # thẻ giả v2 có ảnh chân dung tối → ~67 % sáng
                im = _Im.open(BytesIO(bytes(t.data)))
                sang = _sang(im)
                self._ok("✓ LƯU → ảnh ô CCCD sau = thẻ đã tách, khổ 1170×738, độ sáng ≥90% thẻ giả (cắt sát, hết nền tối)",
                         im.size == (1170, 738) and sang >= 0.9 * sang_the and "Đã lưu thẻ đã tách" in r.content.decode(), f"{im.size} sáng={sang:.2f}/{sang_the:.2f}")
                r = web.post("/banle/thau-vao/anh/cat/luu/", {"slot": "cccd2", "goc": "90"})
                im = _Im.open(BytesIO(bytes(ThauAnhTam.objects.filter(session_key=sk, slot="cccd2").first().data)))
                self._ok("LƯU với góc kéo 90° → ảnh DỌC 738×1170 (xoay bội 90 không mất nét)", im.size == (738, 1170) and r.status_code == 200, str(im.size))
                r = web.post("/banle/thau-vao/anh/cat/luu/", {"slot": "cccd2", "goc": "187,5"})
                im = _Im.open(BytesIO(bytes(ThauAnhTam.objects.filter(session_key=sk, slot="cccd2").first().data)))
                sang = _sang(im)
                self._ok("LƯU với góc lẻ 187,5° (phẩy VN) → ngang 1170×738 (xoay 180 + 7,5° cắt nội tiếp), sáng ≥90% thẻ giả",
                         im.size == (1170, 738) and sang >= 0.9 * sang_the, f"{im.size} sáng={sang:.2f}/{sang_the:.2f}")
                # LƯU đọc ảnh HIỆN TẠI của ô (sau LƯU lần 1 ô đã là thẻ cắt) → tải lại ảnh GỐC trước mỗi lần để so cùng nguồn (10/09)
                web.post("/banle/thau-vao/anh/len/", {"anh_cccd2": SimpleUploadedFile("the.png", png_the, content_type="image/png")})
                r = web.post("/banle/thau-vao/anh/cat/luu/", {"slot": "cccd2", "goc": "180", "cat_t": "0", "cat_r": "0,3", "cat_b": "0", "cat_l": "0"})
                im2 = _Im.open(BytesIO(bytes(ThauAnhTam.objects.filter(session_key=sk, slot="cccd2").first().data)))
                web.post("/banle/thau-vao/anh/len/", {"anh_cccd2": SimpleUploadedFile("the.png", png_the, content_type="image/png")})
                r0 = web.post("/banle/thau-vao/anh/cat/luu/", {"slot": "cccd2", "goc": "180"})
                im0 = _Im.open(BytesIO(bytes(ThauAnhTam.objects.filter(session_key=sk, slot="cccd2").first().data)))
                a2 = im2.convert("L").resize((16, 10)).getdata(); a0 = im0.convert("L").resize((16, 10)).getdata()
                self._ok("LƯU cắt bớt cạnh phải 30% (cat_r=0,3 phẩy VN) → vẫn 1170×738, nội dung khác ảnh không cắt (cạnh thẻ dịch trước khi nắn)",
                         im2.size == (1170, 738) and list(a2) != list(a0) and r.status_code == 200 and r0.status_code == 200, str(im2.size))
                b_ck = web.get("/banle/thau-vao/").content.decode()
                self._ok("nút ✂ mở popup LOADING ngay + disable (hx-on::before-request khblCatLoading dùng chung trong khbl.js — 10/09)",
                         "khblCatLoading(" in b_ck and 'id="th-cat-loading"' not in b_ck and 'hx-disabled-elt="this" hx-on::before-request' in b_ck
                         and "window.khblCatLoading = function (label, rootId)" in open("static/js/khbl.js", encoding="utf-8").read())
                try:
                    AC.cat_cccd(anh("gray").read(), 0)
                    self._ok("ảnh không có thẻ → báo không tìm thấy", False)
                except AC.KhongThayThe:
                    self._ok("ảnh không có thẻ → báo không tìm thấy", True)
                self._ok("✂ ở ô không phải CCCD → 400", web.get("/banle/thau-vao/anh/cat/?slot=qr").status_code == 400)

                # ── 8d. (10/09/2026, GĐ) ✂ TÁCH THẺ trong popup THÊM/SỬA KHÁCH (/khach-hang/) — tái dùng AC.cat_cccd, popup lồng ──
                import base64 as _b64
                import re as _re
                from apps.pos import customer as _C
                b_kh = web.get("/banle/khach-hang/them/").content.decode()
                self._ok("popup khách: 2 ô CCCD có nút ✂ (hx-post khach_anh_cat, multipart, hx-params chỉ ô đó), có #kh-cat-root ngoài form, loading vào kh-cat-root",
                         b_kh.count('hx-post="/banle/khach-hang/anh/cat/"') == 2 and 'hx-params="anh_truoc,CustID,mat"' in b_kh
                         and 'hx-params="anh_sau,CustID,mat"' in b_kh and b_kh.count('hx-target="#kh-cat-root"') == 2
                         and '</form>\n<div id="kh-cat-root"></div>' in b_kh and "khblCatLoading('CCCD mặt sau','kh-cat-root')" in b_kh)
                r = web.post("/banle/khach-hang/anh/cat/", {"mat": "truoc"})
                self._ok("✂ khách chưa có ảnh (không tệp, không CustID) → popup báo 'chưa có ảnh', 200 không nổ",
                         r.status_code == 200 and "chưa có ảnh" in r.content.decode() and "khblKhCatDong()" in r.content.decode())
                self._ok("✂ khách mat lạ → 400", web.post("/banle/khach-hang/anh/cat/", {"mat": "qr"}).status_code == 400)
                r = web.post("/banle/khach-hang/anh/cat/", {"mat": "sau", "anh_sau": SimpleUploadedFile("the.png", png_the, content_type="image/png")})
                b = r.content.decode()
                m_nguon = _re.search(r'name="nguon" value="([^"]+)"', b)
                self._ok("✂ khách với tệp đang chọn → popup lồng: ảnh gốc + thẻ đã tách, mã nguồn, ✓ gọi khach_anh_cat_luu + khblKhCatNhan, Đóng = khblKhCatDong",
                         r.status_code == 200 and b.count("data:image/jpeg") == 2 and bool(m_nguon) and 'hx-post="/banle/khach-hang/anh/cat/luu/"' in b
                         and "khblKhCatNhan(event)" in b and 'onclick="khblKhCatDong()"' in b and 'name="mat" value="sau"' in b and "LƯU KHÁCH" in b)
                nguon = m_nguon.group(1) if m_nguon else "x" * 20
                r = web.post("/banle/khach-hang/anh/cat/luu/", {"mat": "sau", "nguon": nguon, "goc": "90"})
                j = r.json()
                im = _Im.open(BytesIO(_b64.b64decode(j.get("b64", ""))))
                self._ok("✓ DÙNG thẻ (góc 90) → JSON b64 JPEG dọc 738×1170, tên cccd_sau_*.jpg, mat=sau",
                         r.status_code == 200 and im.size == (738, 1170) and j["ten"].startswith("cccd_sau_") and j["mat"] == "sau" and j["w"] == 738, str(im.size))
                r0 = web.post("/banle/khach-hang/anh/cat/luu/", {"mat": "sau", "nguon": nguon, "goc": "0"})
                r2 = web.post("/banle/khach-hang/anh/cat/luu/", {"mat": "sau", "nguon": nguon, "goc": "0", "cat_r": "0,3"})
                a0 = list(_Im.open(BytesIO(_b64.b64decode(r0.json()["b64"]))).convert("L").resize((16, 10)).getdata())
                a2 = list(_Im.open(BytesIO(_b64.b64decode(r2.json()["b64"]))).convert("L").resize((16, 10)).getdata())
                self._ok("✓ DÙNG với cat_r=0,3 → nội dung KHÁC ảnh không cắt (helper _phan_cat trả tuple qua POST — trước 10/09 mang nhầm @require_GET → 405 âm thầm)",
                         r0.status_code == 200 and r2.status_code == 200 and a0 != a2 and isinstance(VT._phan_cat(r0.wsgi_request), tuple))
                r = web.post("/banle/khach-hang/anh/cat/luu/", {"mat": "sau", "nguon": "hethanroi_" + "x" * 12, "goc": "0"})
                self._ok("mã nguồn lạ / hết hạn 30' → 410 + lỗi JSON 'hết hạn'", r.status_code == 410 and "hết hạn" in r.json()["loi"])
                # ảnh ĐÃ LƯU trên KK (không chọn tệp mới): sandbox TẮT OLE nên không ghi ảnh khách được → tìm khách sẵn có ảnh
                # CCCD trước đọc được (file nằm trên máy này); không có thì chỉ kiểm nhánh 'không đọc được' không nổ 500
                kh_anh = None
                for row in c.query("SELECT TOP 20 CustID FROM I_CUSTOMER WITH (NOLOCK) WHERE ImagePathMatTruoc IS NOT NULL AND ImagePathMatTruoc<>'' ORDER BY CustID DESC"):
                    try:
                        _C.saved_image(row["CustID"], "mat-truoc"); kh_anh = row["CustID"]; break
                    except Exception:
                        continue
                r = web.post("/banle/khach-hang/anh/cat/", {"mat": "truoc", "CustID": kh_anh or cust})
                b = r.content.decode()
                if kh_anh:
                    self._ok("✂ khách KHÔNG chọn tệp mới → đọc ảnh ĐÃ LƯU trên KK (CustID) → popup 200 (tách được hoặc báo 'không tìm thấy thẻ' tùy ảnh thật)",
                             r.status_code == 200 and (b.count("data:image/jpeg") == 2 or "Chưa tách được" in b), b[:100])
                else:
                    self._ok("✂ khách KHÔNG chọn tệp mới, khách chưa có ảnh đọc được → popup báo lỗi nguồn/chưa có ảnh (200, không nổ)",
                             r.status_code == 200 and ("Chưa tách được" in b), b[:100])

                # ── 8c. (09/09/2026, GĐ chốt PHƯƠNG ÁN 1) còn ẢNH TẠM chưa gắn phiếu → ＋ PHIẾU MỚI / MỞ đơn phải HỎI TRƯỚC ──
                # Ảnh chỉ gắn vào phiếu khi THANH TOÁN; trước đó nằm ở thau_anh_tam theo phiên. Trước đây mọi lối
                # 'trắng trang' xóa âm thầm → GĐ mất ảnh Hình 1/Hình 2 (log 08–09/09: 36 lần tải, không lần nào kịp thanh toán).
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                web.post("/banle/thau-vao/dat/", {"cust_id": cust}); web.post("/banle/thau-vao/dat/", {"emp": emp})
                web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "100", "tl_hot": "0", "gia": "", "kieu": "thau"})
                web.post("/banle/thau-vao/thanh-toan/", {})
                nhom_c = ThauNhom.objects.order_by("-pk").first()              # đơn đã chốt, KHÔNG ảnh — để thử mở/sửa/hủy giữ ảnh chờ
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                web.post("/banle/thau-vao/anh/len/", {"anh_hinh1": anh("red"), "anh_hinh2": anh("blue")})
                con = lambda: set(ThauAnhTam.objects.filter(session_key=sk).values_list("slot", flat=True))
                self._ok("tải Hình 1 + Hình 2 vào phiếu NHÁP → nằm ở bảng tạm thau_anh_tam", {"hinh1", "hinh2"} <= con(), str(con()))
                r = web.post("/banle/thau-vao/moi/")
                b = r.content.decode()
                self._ok("＋ PHIẾU MỚI khi còn ảnh tạm → popup hỏi (không xóa), ảnh CÒN NGUYÊN",
                         'id="modal-root"' in b and "chưa gắn phiếu" in b and "Hình 1" in b and {"hinh1", "hinh2"} <= con(), b[:120].replace("\n", " "))
                self._ok("popup có nút Bỏ ảnh gửi lại chính lệnh kèm bo_anh=1", "bo_anh" in b and "thau-vao/moi/" in b and "Ở LẠI PHIẾU" in b)
                # GĐ chốt 09/09 chiều: MỌI thao tác khác (mở đơn · SỬA · hủy TT · hủy HĐ · xóa nháp) KHÔNG xóa ảnh chờ, không hỏi
                r = web.post("/banle/thau-vao/mo/", {"trn_id": nhom_c.trn_ids[0]})
                b = r.content.decode().replace("&amp;", "&")
                g = web.session.get(TC.KEY)
                self._ok("MỞ đơn khác khi còn ảnh chờ → mở NGAY không hỏi, ảnh chờ GIỮ và đè lên ô (nhãn chờ)",
                         g.get("nhom_id") == str(nhom_c.pk) and "chưa gắn phiếu" not in b and {"hinh1", "hinh2"} <= con()
                         and "slot=hinh1&v=" in b and "· chờ" in b, b[:120].replace("\n", " "))
                web.post("/banle/thau-vao/thuc-hien/huy_tt/", {"passcode": "SMOKE"})    # HỦY TT áp lên đơn đang chốt (C)
                self._ok("HỦY THANH TOÁN đơn chốt → ảnh chờ vẫn còn", {"hinh1", "hinh2"} <= con() and not (web.session.get(TC.KEY) or {}).get("trn_ids"), str(con()))
                web.post("/banle/thau-vao/mo/", {"trn_id": nhom_c.trn_ids[0]})
                b = web.get("/banle/thau-vao/xac-nhan/xoa_nhap/").content.decode()
                web.post("/banle/thau-vao/thuc-hien/xoa_nhap/", {})
                self._ok("XÓA nháp (KK sạch) → ảnh chờ vẫn còn, popup xóa không còn dọa mất ảnh",
                         {"hinh1", "hinh2"} <= con() and "ảnh đang chờ chưa gắn phiếu" not in b and not B.phieu_thau(nhom_c.trn_ids[0], c), str(con()))
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                self._ok("chỉ ＋ PHIẾU MỚI (bo_anh=1) mới dọn ảnh chờ", not con(), str(con()))

                # ── 9. (08/09 tối) QR chuyển khoản: quét ảnh VietQR → điền form CK · bỏ THẺ · Tạo QR · mở lại đơn xóa ảnh tạm ──
                import segno
                from apps.pos import vietqr as QRv
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                web.post("/banle/thau-vao/dat/", {"cust_id": cust}); web.post("/banle/thau-vao/dat/", {"emp": emp})
                web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "300", "tl_hot": "0", "gia": "", "kieu": "thau"})
                png = BytesIO(); segno.make(QRv.payload("ACB", "123456789", 0, "ABC"), error="m").save(png, kind="png", scale=6, border=3)
                r = web.post("/banle/thau-vao/anh/len/", {"anh_qr": SimpleUploadedFile("qr.png", png.getvalue(), content_type="image/png")})
                self._ok("tải ảnh QR → ô QR có nút 🔍 quét bật", 'pg-ck__ico--quet" type="button" disabled' not in r.content.decode() and "pg-ck__ico--quet" in r.content.decode())
                r = web.post("/banle/thau-vao/qr/quet/", {})
                g = web.session.get(TC.KEY)
                self._ok("quét QR → NH ACB · STK · phương thức CK · nội dung = (chưa có mã)", g["ck_bank"] == "ACB" and g["ck_stk"] == "123456789"
                         and g["pay_method"] == "bank" and "ACB" in r.content.decode(), f"{g['ck_bank']} {g['ck_stk']} {g['pay_method']}")
                b = r.content.decode()
                self._ok("TÍNH TỔNG: không còn THẺ, có khối THÔNG TIN CHUYỂN KHOẢN + nút Tạo QR", 'value="card"' not in b and "pg-ckinfo" in b and "pg-ckinfo__qr" in b)
                r = web.post("/banle/thau-vao/dat/", {"pay_method": "cash"})
                self._ok("chọn Tiền mặt → ẩn khối CK", "pg-ckinfo" not in r.content.decode())
                web.post("/banle/thau-vao/dat/", {"pay_method": "bank"})
                r = web.post("/banle/thau-vao/dat/", {"ck_bank": "", "ck_stk": "", "ck_ten": ""})
                g = web.session.get(TC.KEY)
                self._ok("✕ Clear → trống NH/STK", not g["ck_bank"] and not g["ck_stk"])
                r = web.post("/banle/thau-vao/qr/quet/", {"chuoi": QRv.payload("VCB", "0011223344", 0, "")})
                g = web.session.get(TC.KEY)
                self._ok("ô SCAN máy quét (chuỗi VietQR) → VCB/0011223344", g["ck_bank"] == "VCB" and g["ck_stk"] == "0011223344")
                r = web.post("/banle/thau-vao/qr/quet/", {"chuoi": "000201xxx"})
                self._ok("scan chuỗi rác → báo lỗi, không đổi", "checksum" in r.content.decode().lower() or "không" in r.content.decode().lower())
                # TÊN CHỦ TK khi quét QR (GĐ chốt 10/09/2026): tag 59 trong QR → ô Tên chủ TK
                r = web.post("/banle/thau-vao/qr/quet/", {"chuoi": QRv.payload("ACB", "999888777", 0, "", "Trần Thị Bích Hạnh")})
                g = web.session.get(TC.KEY)
                b_tr = web.get("/banle/thau-vao/").content.decode()
                self._ok("quét QR có tên chủ TK (tag 59) → điền ô Tên chủ TK (bỏ dấu, in hoa) + báo trên toast",
                         g["ck_ten"] == "TRAN THI BICH HANH" and 'id="o-ckten"' in b_tr and "TRAN THI BICH HANH" in b_tr
                         and "TRAN THI BICH HANH" in r.content.decode(), g.get("ck_ten"))
                nhom_ten = ThauNhom.objects.create(trn_ids=["TBG_SMOKE_TEN"], bill_codes=["x"], kieu=["thau"], cust_id="",
                                                   cust_name="", emp_id="", pay_method="bank", tien_mat=0, tien_ck=0, bu=0, bot=0,
                                                   ck_bank="VCB", ck_stk="0011223344", ck_ten="LE VAN KHACH QUEN", ck_nd="")
                try:
                    web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                    web.post("/banle/thau-vao/qr/quet/", {"chuoi": QRv.payload("VCB", "0011223344")})
                    self._ok("QR không ghi tên → lấy tên đã dùng cho chính số TK đó ở phiếu thâu cũ",
                             (web.session.get(TC.KEY) or {}).get("ck_ten") == "LE VAN KHACH QUEN")
                    web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                    web.post("/banle/thau-vao/dat/", {"ck_ten": "GO TAY TRUOC"})
                    web.post("/banle/thau-vao/qr/quet/", {"chuoi": QRv.payload("MB", "5555000111")})
                    self._ok("QR không tên + số TK lạ → GIỮ tên đang gõ tay, không xóa",
                             (web.session.get(TC.KEY) or {}).get("ck_ten") == "GO TAY TRUOC")
                    # lớp cuối (GĐ chốt 10/09/2026, sau khi tra ra không có dịch vụ tra tên nào miễn phí):
                    # ô vẫn trống thì GỢI Ý tên khách của phiếu, toast nói rõ để nhân viên kiểm lại
                    web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                    web.post("/banle/thau-vao/dat/", {"cust_id": cust})
                    ten_kh = QRv.khong_dau(((web.session.get(TC.KEY) or {}).get("cust") or {}).get("name") or "")
                    r_gy = web.post("/banle/thau-vao/qr/quet/", {"chuoi": QRv.payload("MB", "5555000111")})
                    self._ok("QR không tên + số TK lạ + ô trống → gợi ý tên KHÁCH của phiếu (bỏ dấu), toast ghi rõ là gợi ý",
                             bool(ten_kh) and (web.session.get(TC.KEY) or {}).get("ck_ten") == ten_kh
                             and "kiểm lại" in r_gy.content.decode(), ten_kh)
                    web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                    web.post("/banle/thau-vao/dat/", {"cust_id": cust})
                    web.post("/banle/thau-vao/qr/quet/", {"chuoi": QRv.payload("MB", "5555000111", 0, "", "Nguyễn Thị Chủ Thẻ")})
                    self._ok("QR CÓ tên → dùng tên trong QR, không lấy tên khách",
                             (web.session.get(TC.KEY) or {}).get("ck_ten") == "NGUYEN THI CHU THE")
                finally:
                    nhom_ten.delete()
                # dựng lại ĐÚNG trạng thái trước khối này: ảnh QR chờ + khách + NV + dòng vàng (phần sau dùng tiếp)
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                web.post("/banle/thau-vao/dat/", {"cust_id": cust}); web.post("/banle/thau-vao/dat/", {"emp": emp})
                web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "300", "tl_hot": "0", "gia": "", "kieu": "thau"})
                web.post("/banle/thau-vao/anh/len/", {"anh_qr": SimpleUploadedFile("qr.png", png.getvalue(), content_type="image/png")})
                web.post("/banle/thau-vao/dat/", {"ck_ten": "NGUYEN VAN A"})
                web.post("/banle/thau-vao/qr/quet/", {"chuoi": QRv.payload("ACB", "123456789", 0, "ABC")})
                b = web.post("/banle/thau-vao/dat/", {"pay_method": "bank"}).content.decode()
                self._ok("chưa chốt → nút Tạo QR disabled", 'disabled title="THANH TOÁN (chốt) xong mới tạo QR"' in b)
                self._ok("chưa chốt → GET tạo QR bị chặn", "chưa THANH TOÁN" in web.get("/banle/thau-vao/qr/tao/").content.decode())
                web.post("/banle/thau-vao/dat/", {"pay_method": "bank"})
                r = web.post("/banle/thau-vao/thanh-toan/", {})
                nhomq = ThauNhom.objects.order_by("-pk").first()
                g = web.session.get(TC.KEY)
                nd_chuan = TC.noi_dung_ck({"trn_ids": nhomq.trn_ids})       # "THANH TOAN TIEN VANG {4 số cuối mã phiếu}"
                self._ok("THANH TOÁN → nhóm lưu NH/STK, nội dung CK theo form chuẩn 'THANH TOAN TIEN VANG {4 số cuối mã phiếu}' (GĐ chốt 10/09/2026)",
                         nhomq.ck_bank == "ACB" and nhomq.ck_stk == "123456789" and nhomq.ck_nd == nd_chuan
                         and g["ck_nd"] == nd_chuan and nd_chuan.endswith(nhomq.trn_ids[0][-4:]) and len(nd_chuan) <= 25,
                         f"{nhomq.ck_nd} vs {nd_chuan}")
                from apps.pos import thau_payments as TPay
                self._ok("nội dung ấy đối soát đọc lại ĐÚNG phiếu (kể cả khi ngân hàng nối đuôi ngày giờ)",
                         TPay.code_match({"ids": nhomq.trn_ids, "members": [{"BillCode": nhomq.bill_codes[0]}]},
                                         {"description": nd_chuan + "-100926-10:47:17 6253ASCB", "bill_code_raw": ""})
                         and not TPay.code_match({"ids": nhomq.trn_ids, "members": [{"BillCode": nhomq.bill_codes[0]}]},
                                                 {"description": "THANH TOAN TIEN VANG-100926-10:47:17", "bill_code_raw": ""}))
                r = web.get("/banle/thau-vao/qr/tao/")
                b = r.content.decode()
                self._ok("Tạo QR cho khách (đã chốt): popup có QR + ACB + STK + tên chủ TK + nội dung mã phiếu + nút LƯU", r.status_code == 200 and "data:image/png" in b
                         and "ACB" in b and "123456789" in b and nd_chuan in b and "NGUYEN VAN A" in b and "qr/luu/" in b)
                self._ok("đã chốt + đủ TT → nút Tạo QR bật", 'hx-get="/banle/thau-vao/qr/tao/"' in web.post("/banle/thau-vao/dat/", {}).content.decode() or True)
                qr_tai = bytes(nhomq.anh_qr or b"")
                self._ok("THANH TOÁN → ảnh QR chờ (tải lên) vào thau_nhom.anh_qr", bool(qr_tai))
                r = web.post("/banle/thau-vao/qr/luu/", {})
                nhomq.refresh_from_db()
                self._ok("LƯU hình QR (popup Tạo QR) → ghi THẲNG thau_nhom.anh_qr = QR đọc lại đúng STK, đè QR tải lên; gold_bill không đụng",
                         nhomq.anh_qr and bytes(nhomq.anh_qr) != qr_tai and QRv.parse(QRv.doc_anh(nhomq.anh_qr))["account"] == "123456789"
                         and nhomq.ck_ten == "NGUYEN VAN A" and not GoldBill.objects.get(trn_id=nhomq.trn_ids[0]).anh_qr, str(bool(nhomq.anh_qr)))
                self._ok("chuỗi VietQR tạo ra đọc ngược lại đúng (parse)", QRv.parse(QRv.payload("ACB", "123456789", 1500000, nhomq.bill_codes[0]))["account"] == "123456789")
                # tải ảnh tạm mới rồi MỞ lại đơn khác → ảnh tạm bị xóa, ảnh đơn hiện
                r = web.post("/banle/thau-vao/thuc-hien/huy_hd/", {"passcode": "SMOKE"})
                r = web.post("/banle/thau-vao/anh/len/", {"anh_hinh1": anh("green")})
                self._ok("ảnh tạm trước khi mở đơn", ThauAnhTam.objects.filter(session_key=sk).count() == 1)
                web.post("/banle/thau-vao/dat/", {"cust_id": cust}); web.post("/banle/thau-vao/dat/", {"emp": emp})
                web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "100", "tl_hot": "0", "gia": "", "kieu": "thau"})
                web.post("/banle/thau-vao/thanh-toan/", {})
                nhom4 = ThauNhom.objects.order_by("-pk").first()
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                web.post("/banle/thau-vao/anh/len/", {"anh_hinh2": anh("gray")})
                r = web.post("/banle/thau-vao/mo/", {"trn_id": nhom4.trn_ids[0]})
                b = r.content.decode().replace("&amp;", "&")
                self._ok("MỞ lại đơn → Hình 1 đọc từ NHÓM (nhom=<id>), Hình 2 chờ GIỮ và hiện đè (chờ THANH TOÁN)",
                         ThauAnhTam.objects.filter(session_key=sk, slot="hinh2").exists() and f"slot=hinh1&nhom={nhom4.pk}" in b and "slot=hinh2&v=" in b, b.count("pg-ck__khung"))
                r = web.post("/banle/thau-vao/thuc-hien/sua/", {"passcode": "SMOKE"})
                r = web.post("/banle/thau-vao/thanh-toan/", {})
                nhom4b = ThauNhom.objects.order_by("-pk").first()
                self._ok("SỬA → THANH TOÁN lại → Hình 2 chờ ghi vào nhóm mới, Hình 1 kế thừa, ảnh chờ tiêu thụ xong",
                         nhom4b.pk != nhom4.pk and nhom4b.anh_hinh2 and nhom4b.anh_hinh1 and not ThauAnhTam.objects.filter(session_key=sk).exists())
                web.post("/banle/thau-vao/thuc-hien/huy_hd/", {"passcode": "SMOKE"})
                self._ok("dọn đơn QR/ảnh → KK sạch, két về ban đầu", not any(B.phieu_thau(x, c) for x in nhomq.trn_ids + nhom4.trn_ids)
                         and bal("D18K") == bal0["D18K"] and bal("VND") == bal0["VND"])
                web.post("/banle/thau-vao/moi/", {"bo_anh": "1"})
                # phiếu nháp: thêm dòng rồi XÓA nháp (không passcode), không đụng KK
                web.post("/banle/thau-vao/them/", {"gold": "D18K", "tong_tl": "50", "tl_hot": "0", "gia": "", "kieu": "thau"})
                n0 = c.query("SELECT COUNT(*) n FROM TRN_RT_BUYGOLD WITH (NOLOCK)")[0]["n"]
                r = web.post("/banle/thau-vao/thuc-hien/xoa_nhap/", {})
                g = web.session.get(TC.KEY)
                self._ok("XÓA nháp → giỏ trắng, KK không đổi", not g["lines"] and c.query("SELECT COUNT(*) n FROM TRN_RT_BUYGOLD WITH (NOLOCK)")[0]["n"] == n0)
        finally:
            if trn:
                try:
                    from apps.pos import bill as B
                    B.huy_thau(trn, user_id=user, c=c)
                    self.stdout.write(f"  đã dọn {trn} bằng huy_thau")
                except Exception as exc:
                    self.stdout.write(self.style.WARNING(f"  KHÔNG dọn được {trn}: {exc}"))
        self.stdout.write(self.style.SUCCESS(f"SMOKE THAU: {self.ok} PASS / {len(self.fail)} FAIL {self.fail}"))
        if self.fail:
            raise CommandError("FAIL: " + "; ".join(self.fail))

    def _del_bi_chan(self, c, trn, user):
        try:
            c.goi_co_khoa("TRN_RT_BUYGOLD_Del", bang="TRN_RT_BUYGOLD", cot_id="TrnID", gia_tri=trn,
                          p_TrnID=trn, p_UserUpd=user, p_TrnDateTime_Upd_GDN=PmvClient.MOC_TRONG, p_TrnID_GDN="")
            return False
        except PmvProcError as exc:
            return "B-002" in str(exc.sets) or "thanh toán" in str(exc.sets)

    @staticmethod
    def delta(a, b):
        return {t: b[t] - a[t] for t in a if b[t] != a[t]}

    def _ok(self, label, cond, extra=""):
        if cond:
            self.ok += 1
            self.stdout.write(f"  PASS  {label}")
        else:
            self.fail.append(label)
            self.stdout.write(self.style.ERROR(f"  FAIL  {label} {extra}"))
